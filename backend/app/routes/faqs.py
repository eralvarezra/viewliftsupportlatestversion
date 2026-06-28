# backend/app/routes/faqs.py
"""FAQ management routes for SCHN+ Support Assistant."""
import os
import tempfile
from typing import List

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Query
from sqlalchemy.orm import Session

from app.auth.routes import get_current_user, require_admin
from app.database import get_db
from app.models import User, FAQDocument, FAQChunk
from app.schemas import FAQDocumentResponse, FAQChunkResponse
from app.services.docx_processor import DocxProcessor
from app.services.xlsx_processor import XlsxProcessor
from app.services.local_embeddings import LocalEmbeddingService

def _is_liv_golf_audit(filepath: str) -> bool:
    """Return True if the XLSX has both 'CMS Export' and 'Tune In By Country' sheets."""
    from openpyxl import load_workbook as _load
    try:
        wb = _load(filepath, read_only=True)
        try:
            names = wb.sheetnames
            return "CMS Export" in names and "Tune In By Country" in names
        finally:
            wb.close()
    except Exception:
        return False


router = APIRouter()

# Supported file extensions
SUPPORTED_EXTENSIONS = {'.docx', '.xlsx'}


@router.get("/", response_model=List[FAQDocumentResponse])
async def list_faqs(
    platform_id: int = Query(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """List all FAQ documents for a given platform.

    Any authenticated user can view the list of uploaded FAQ documents.

    Args:
        platform_id: The platform to filter by.
        current_user: The authenticated user.
        db: Database session.

    Returns:
        List of FAQ documents with their metadata.
    """
    documents = (
        db.query(FAQDocument)
        .filter(FAQDocument.platform_id == platform_id)
        .order_by(FAQDocument.uploaded_at.desc())
        .all()
    )
    return documents


@router.post("/upload", response_model=FAQDocumentResponse)
async def upload_faq(
    file: UploadFile = File(...),
    platform_id: int = Form(...),
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    """Upload a DOCX or XLSX file as FAQ document.

    Admin only endpoint. Accepts a DOCX or XLSX file, processes it into chunks,
    generates embeddings for each chunk, and saves to the database.

    Args:
        file: The uploaded file.
        current_user: The authenticated admin user.
        db: Database session.

    Returns:
        The created FAQ document metadata.

    Raises:
        HTTPException: If file is invalid or processing fails.
    """
    # Validate file type
    if not file.filename:
        raise HTTPException(status_code=400, detail="No filename provided")

    file_ext = os.path.splitext(file.filename)[1].lower()
    if file_ext not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type. Supported types: {', '.join(SUPPORTED_EXTENSIONS)}"
        )

    # Save uploaded file temporarily
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=file_ext) as tmp_file:
            content = await file.read()
            tmp_file.write(content)
            tmp_path = tmp_file.name
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to save uploaded file: {str(e)}"
        )

    try:
        # Process based on file type
        if file_ext == '.docx':
            processor = DocxProcessor()
            result = processor.process_docx(tmp_path)
        elif file_ext == '.xlsx':
            if _is_liv_golf_audit(tmp_path):
                from app.services.liv_golf_processor import LivGolfAuditProcessor
                result = LivGolfAuditProcessor().process(tmp_path)
            else:
                from app.services.app_store_links_processor import AppStoreLinksProcessor
                if AppStoreLinksProcessor.detect(tmp_path):
                    result = AppStoreLinksProcessor().process(tmp_path)
                else:
                    processor = XlsxProcessor()
                    result = processor.process_xlsx(tmp_path)
        else:
            raise HTTPException(status_code=400, detail="Unsupported file type")
    except FileNotFoundError:
        raise HTTPException(status_code=400, detail="File not found")
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to process file: {str(e)}"
        )
    finally:
        # Clean up temp file (may fail on Windows if file is still locked)
        try:
            if os.path.exists(tmp_path):
                os.unlink(tmp_path)
        except PermissionError:
            # File still locked on Windows, will be cleaned up later
            pass

    # Detect document type from filename
    filename_lower = file.filename.lower()
    location_kw = ("location_rule", "travel_rule", "geo_rule", "location-rule")
    zip_kw = ("zip", "zipcode", "zip_code", "cobertura")
    doc_type = "location_rules" if any(k in filename_lower for k in location_kw) else ("zipcode" if any(k in filename_lower for k in zip_kw) else "faq")

    # Create FAQ document record
    faq_document = FAQDocument(
        filename=file.filename,
        uploaded_by=current_user.id,
        chunk_count=result["total_chunks"],
        document_type=doc_type,
        platform_id=platform_id,
    )
    db.add(faq_document)
    db.commit()
    db.refresh(faq_document)

    # Generate embeddings and create chunks
    embedding_service = LocalEmbeddingService()

    try:
        for idx, chunk_content in enumerate(result["chunks"]):
            # Generate embedding for the chunk
            embedding = embedding_service.get_embedding(chunk_content)
            embedding_bytes = embedding_service.serialize_embedding(embedding)

            # Create chunk record
            faq_chunk = FAQChunk(
                document_id=faq_document.id,
                content=chunk_content,
                embedding=embedding_bytes,
                chunk_metadata={"chunk_index": idx},
                platform_id=platform_id,
            )
            db.add(faq_chunk)

        db.commit()
    except Exception as e:
        # Rollback document creation if embedding fails
        db.delete(faq_document)
        db.commit()
        raise HTTPException(
            status_code=500,
            detail=f"Failed to generate embeddings: {str(e)}"
        )

    return faq_document


@router.delete("/{faq_id}")
async def delete_faq(
    faq_id: int,
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    """Delete an FAQ document.

    Admin only endpoint. Deletes an FAQ document and all its associated
    chunks (cascade delete).

    Args:
        faq_id: The ID of the FAQ document to delete.
        current_user: The authenticated admin user.
        db: Database session.

    Returns:
        Success message.

    Raises:
        HTTPException: If document not found.
    """
    faq_document = db.query(FAQDocument).filter(FAQDocument.id == faq_id).first()

    if not faq_document:
        raise HTTPException(status_code=404, detail="FAQ document not found")

    # Delete the document (chunks will be cascade deleted)
    db.delete(faq_document)
    db.commit()

    return {"message": "FAQ document deleted successfully"}


@router.get("/{faq_id}/chunks", response_model=List[FAQChunkResponse])
async def get_faq_chunks(
    faq_id: int,
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    """Return all chunks for a specific FAQ document. Admin only."""
    document = db.query(FAQDocument).filter(FAQDocument.id == faq_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="FAQ document not found")

    chunks = db.query(FAQChunk).filter(FAQChunk.document_id == faq_id).all()
    chunks.sort(
        key=lambda c: c.chunk_metadata.get("chunk_index", 0) if c.chunk_metadata else 0
    )

    return [
        FAQChunkResponse(
            id=chunk.id,
            content=chunk.content,
            chunk_index=chunk.chunk_metadata.get("chunk_index", idx) if chunk.chunk_metadata else idx,
        )
        for idx, chunk in enumerate(chunks)
    ]